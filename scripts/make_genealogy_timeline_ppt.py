# -*- coding: utf-8 -*-
"""예수님 족보(마 1장) 3기 마일스톤 계단식 연표 — 1장짜리 PPT."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
import sys

OUT = sys.argv[1]
FONT = "맑은 고딕"

# ---------- data: (기, 이름, 시작, 끝, 추정여부, 정보) — BC는 음수 ----------
ROWS = [
    # 1기 — 약속의 시대 (생애)
    (1, "아브라함", -2166, -1991, False, "2166–1991 · 부르심(2091) · 횃불 언약(창 15)"),
    (1, "이삭", -2066, -1886, False, "2066–1886 · 모리아산 번제(창 22)"),
    (1, "야곱", -2006, -1859, False, "2006–1859 · '이스라엘' · 애굽 이주(1876)"),
    (1, "유다 👩다말", -1916, -1800, True, "약 1916– · 다말(창 38) — 족보 첫 여인"),
    (1, "베레스", -1880, -1780, True, "약 1880– · 다말의 쌍둥이(창 38:29)"),
    (1, "헤스론", -1860, -1760, True, "애굽 이주 명단(창 46:12)"),
    (1, "람", -1760, -1560, True, "애굽 체류기 · 400년 종살이 (연대 불명)"),
    (1, "아미나답", -1530, -1450, True, "약 1500 · 딸이 아론의 아내(출 6:23)"),
    (1, "나손", -1490, -1407, True, "출애굽(1446) · 유다 지파 족장(민 1:7)"),
    (1, "살몬 👩라합", -1440, -1360, True, "가나안 입성(1406) · 라합"),
    (1, "보아스 👩룻", -1160, -1080, True, "사사 시대 · 모압 여인 룻(룻 2–4)"),
    (1, "오벳", -1120, -1040, True, "나오미의 품에서(룻 4:17)"),
    (1, "이새", -1100, -1010, True, "베들레헴 사람(삼상 16)"),
    (1, "다윗", -1040, -970, False, "1040–970 · 즉위 1010 · 다윗 언약(삼하 7)"),
    # 2기 — 왕국과 몰락 (재위)
    (2, "솔로몬 👩밧세바", -970, -931, False, "재위 970–931 · 성전 완공(959)"),
    (2, "르호보암 ❌", -931, -913, False, "931–913 · 왕국 분열(931)"),
    (2, "아비야 ❌", -913, -911, False, "913–911"),
    (2, "아사 ✅", -911, -870, False, "911–870 · 우상 제거 개혁"),
    (2, "여호사밧 ✅", -873, -848, False, "873–848 · 아합 집안과 혼인 동맹"),
    (2, "요람 ❌", -853, -841, False, "853–841 · 아달랴와 결혼 (이후 3왕 생략: 아하시야·요아스·아마샤)"),
    (2, "웃시야 ✅", -792, -740, False, "792–740 · 그가 죽던 해 이사야 소명"),
    (2, "요담 ✅", -750, -732, False, "750–732"),
    (2, "아하스 ❌", -735, -715, False, "735–715 · 재위 중 북이스라엘 멸망(722)"),
    (2, "히스기야 ✅", -715, -686, False, "715–686 · 산헤립 포위와 구원(701)"),
    (2, "므낫세 ❌", -697, -642, False, "697–642 · 성전에 우상, 멸망의 원인"),
    (2, "아몬 ❌", -642, -640, False, "642–640"),
    (2, "요시야 ✅", -640, -609, False, "640–609 · 율법책 발견·대개혁(622) (이후 여호아하스·여호야김 생략)"),
    (2, "여고냐와 형제들 ❌", -598, -586, False, "598–597 재위 · 포로(597) · 형제 시드기야 때 성전 불탐(586)"),
    # 3기 — 회복의 시대
    (3, "여고냐", -597, -561, False, "포로 597 → 석방 561 · 왕의 식탁(왕하 25:27)"),
    (3, "스알디엘", -595, -530, True, "바벨론에서 출생(대상 3:17)"),
    (3, "스룹바벨", -570, -500, True, "귀환 538 · 성전 재건 516 · 인장 반지(학 2:23)"),
    (3, "아비훗", -530, -470, True, "↓ 여기서부터 요셉까지 구약 기록 없음 (연대 추정)"),
    (3, "엘리아김", -480, -420, True, "페르시아 시대 · 에스더·에스라·느헤미야·말라기"),
    (3, "아소르", -430, -370, True, "말라기 이후 '침묵의 400년'"),
    (3, "사독", -380, -320, True, "알렉산더 대왕(333)"),
    (3, "아킴", -330, -270, True, "헬라 시대"),
    (3, "엘리웃", -280, -220, True, "프톨레미 지배"),
    (3, "엘르아살", -230, -160, True, "셀류시드 · 성전 모독(167)·마카비 정화(164)"),
    (3, "맛단", -170, -100, True, "하스몬 왕조"),
    (3, "야곱", -110, -40, True, "로마 점령(63)"),
    (3, "요셉 👩마리아", -50, 10, True, "헤롯 대왕(37–4) · 다윗의 자손(마 1:20)"),
    (3, "예수 그리스도", -5, 30, False, "탄생 약 BC 6–4 · 십자가 약 AD 30–33"),
]

MILESTONES = [
    (-2091, "부르심"), (-1876, "애굽 이주"), (-1446, "출애굽"), (-1406, "가나안"),
    (-1010, "다윗 즉위"), (-959, "성전 완공"), (-931, "분열"), (-722, "북 멸망"),
    (-586, "성전 불탐"), (-538, "귀환"), (-516, "성전 재건"), (-164, "성전 정화"),
    (-5, "예수 탄생"),
]
HIGHLIGHT = -586

PCOL = {1: RGBColor(0xF2, 0xC1, 0x4E), 2: RGBColor(0x4F, 0xA3, 0xE0), 3: RGBColor(0x5C, 0xC9, 0x8A)}
PBAND = {1: RGBColor(0x2A, 0x25, 0x16), 2: RGBColor(0x16, 0x22, 0x2E), 3: RGBColor(0x17, 0x28, 0x1E)}
PNAME = {1: "1기 · 약속 (아브라함 → 다윗, 14대)",
         2: "2기 · 왕국과 몰락 (다윗 → 바벨론 포로, 14대)",
         3: "3기 · 회복 (바벨론 포로 → 그리스도, 14대)"}
BG = RGBColor(0x14, 0x14, 0x14)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x9A, 0x9A, 0x9A)
DIM = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xFF, 0x4D, 0x4D)

# ---------- geometry ----------
SW, SH = 13.333, 7.5
X0, X1 = 1.35, SW - 0.2            # plot area (left leaves room for period labels)
T0, T1 = -2200, 40
TOP = 1.02                          # first row top
BOTTOM = SH - 0.28
N = len(ROWS)
PITCH = (BOTTOM - TOP) / N
BAR_H = 0.042
TXT_H = 0.09


def xpos(year):
    return X0 + (year - T0) / (T1 - T0) * (X1 - X0)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG


def box(x, y, w, h, fill=None, line=None, dash=False, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = MSO_LINE.DASH
    sp.shadow.inherit = False
    return sp


def text(x, y, w, h, t, size, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = t
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1], {"typeface": FONT})
        rPr.append(el)
    return tb


def vline(x, y0, y1, color, w=0.5, dash=True):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y0), Inches(x), Inches(y1))
    ln.line.color.rgb = color; ln.line.width = Pt(w)
    if dash:
        ln.line.dash_style = MSO_LINE.DASH
    return ln


# ---------- title ----------
t = box(0.25, 0.12, 5.6, 0.42, fill=RGBColor(0x8B, 0x1E, 0x1E), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(0.4, 0.12, 5.4, 0.42, "예수 그리스도의 족보 — 마태복음 1:1-17 (3 × 14대)", 16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(6.05, 0.16, 7.1, 0.2, "네모: 왼쪽 변 = 시작, 오른쪽 변 = 끝 · 왕은 재위 기간, 그 외는 생애 · 점선 = 연대 추정(성경에 연대 기록 없음)", 8, GREY)
text(6.05, 0.36, 7.1, 0.2, "✅ 선한 왕  ❌ 악한 왕  👩 마태가 넣은 여인 · 연대: 보수적 주류 연대기 근사값(출애굽 1446 · 다윗 즉위 1010 · 분열 931)", 8, GREY)

# ---------- period bands + labels ----------
for p in (1, 2, 3):
    idx = [i for i, r in enumerate(ROWS) if r[0] == p]
    y0 = TOP + idx[0] * PITCH - 0.02
    y1 = TOP + (idx[-1] + 1) * PITCH
    box(0.2, y0, SW - 0.35, y1 - y0, fill=PBAND[p])
    box(0.2, y0, 0.07, y1 - y0, fill=PCOL[p])
    text(0.32, y0 + 0.04, 1.0, 0.2, f"{p}기", 13, PCOL[p], bold=True)
    lbl = {1: "약속", 2: "왕국과 몰락", 3: "회복"}[p]
    text(0.32, y0 + 0.30, 1.0, 0.2, lbl, 9, PCOL[p], bold=True)
    sub = {1: "아브라함 → 다윗", 2: "다윗 → 바벨론 포로", 3: "포로 → 그리스도"}[p]
    text(0.32, y0 + 0.48, 1.0, 0.2, sub, 7, GREY)

# ---------- time axis + milestones ----------
AX_Y = 0.86
ax = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X0), Inches(AX_Y), Inches(X1), Inches(AX_Y))
ax.line.color.rgb = GREY; ax.line.width = Pt(0.75)
for yr in range(-2200, 1, 100):
    x = xpos(yr)
    vline(x, AX_Y, BOTTOM, RGBColor(0x2E, 0x2E, 0x2E), 0.4)
    if yr % 200 == 0:
        lab = f"BC {-yr}" if yr < 0 else "BC/AD"
        text(x - 0.3, AX_Y + 0.02, 0.6, 0.12, lab, 6, GREY, align=PP_ALIGN.CENTER)

for i, (yr, lab) in enumerate(MILESTONES):
    x = xpos(yr)
    col = RED if yr == HIGHLIGHT else WHITE
    d = box(x - 0.045, AX_Y - 0.045, 0.09, 0.09, fill=col, shape=MSO_SHAPE.DIAMOND)
    vline(x, AX_Y, BOTTOM, RED if yr == HIGHLIGHT else DIM, 1.25 if yr == HIGHLIGHT else 0.4)
    ty = AX_Y - (0.31, 0.22, 0.13)[i % 3]
    text(x - 0.45, ty, 0.9, 0.11, f"{lab} {-yr if yr < 0 else ''}".strip(), 6, col, bold=(yr == HIGHLIGHT), align=PP_ALIGN.CENTER)

# ---------- rows (staircase) ----------
CHAR_W = 0.052  # approx width per char @6pt (Korean)
for i, (p, name, a, b, est, info) in enumerate(ROWS):
    y = TOP + i * PITCH
    xa, xb = xpos(a), xpos(b)
    w = max(xb - xa, 0.05)
    if est:
        box(xa, y, w, BAR_H, fill=None, line=PCOL[p], dash=True, lw=0.75)
    else:
        box(xa, y, w, BAR_H, fill=PCOL[p])
    label = f"{name}  ·  {info}"
    tw = len(label) * CHAR_W + 0.1
    ty = y + BAR_H + 0.004
    if xa + tw <= SW - 0.2:
        text(xa, ty, tw, TXT_H, label, 5.5, WHITE)
    else:  # near the right edge → right-align to bar end
        right = min(xa + w, SW - 0.2)
        text(right - tw, ty, tw, TXT_H, label, 5.5, WHITE, align=PP_ALIGN.RIGHT)

# ---------- 586 callout (left of the red line, in the empty 3기 area) ----------
x = xpos(HIGHLIGHT)
iy = [i for i, r in enumerate(ROWS) if r[1] == "아비훗"][0]
cy = TOP + iy * PITCH
text(x - 4.1, cy, 4.0, 0.14, "10/18 설교 본문 (왕하 25) — 성전이 불탄 날 ▶", 7, RED, bold=True, align=PP_ALIGN.RIGHT)
text(x - 4.1, cy + 0.16, 4.0, 0.14, "족보의 두 번째 경첩 '바벨론으로 사로잡혀 갈 때' (마 1:11, 12, 17)", 6, RED, align=PP_ALIGN.RIGHT)

prs.save(OUT)
print("saved", OUT, "rows", N, "pitch", round(PITCH, 3))
