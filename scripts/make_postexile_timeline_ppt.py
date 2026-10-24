# -*- coding: utf-8 -*-
"""포로기와 귀환 — 왕들과 선지자들 마일스톤 연표 (BC 630–420), 대형 1장 + 근거 메모 1장.

- 왕: 즉위~폐위를 네모의 왼쪽/오른쪽 변으로 (애굽 · 유다 왕·총독 · 바벨론 · 공동통치/논쟁 · 메대·바사)
- 선지자·말씀 지도자: 성경 근거 활동 기간을 네모로, 번호 → 아래 카드(활동·대상·핵심 메시지·갈등·대표 본문·하나님의 의도)
- 점선 = 연대 추정/논쟁. 연대는 관례적 연대(바벨론 연대기·페르시아 왕명록, 예루살렘 함락 586년설).

사용: python scripts/make_postexile_timeline_ppt.py <출력.pptx>
"""
import math
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE

OUT = sys.argv[1]
FONT = "맑은 고딕"
RGB = lambda h: RGBColor.from_string(h)
BG, WHITE, GREY, DIM, CARD = RGB("141414"), RGB("FFFFFF"), RGB("A0A0A0"), RGB("4A4A4A"), RGB("1F1F1F")
INK = RGB("141414")
COL = {
    "egypt": RGB("C9A86A"), "judah": RGB("4FA3E0"), "gov": RGB("8CC4EC"), "bab": RGB("E8955F"),
    "co": RGB("D9B07A"), "persia": RGB("E0C04F"), "prophet": RGB("C78BE8"), "leader": RGB("5CC98A"),
    "judg": RGB("FF4D4D"), "ret": RGB("5CC98A"), "word": RGB("F2D16B"), "pol": RGB("A0A0A0"),
}

# ── 시간축 ──────────────────────────────────────────────
Y_START, Y_END = 630, 420          # BC
SW = 56.0                          # PowerPoint 최대 폭(인치)
LBL, RM = 3.4, 0.4
W = SW - LBL - RM
PER_YEAR = W / (Y_START - Y_END)


def X(bc):
    """BC 연도 → x(인치). 범위 밖은 잘라낸다."""
    bc = min(max(bc, Y_END), Y_START)
    return LBL + (Y_START - bc) * PER_YEAR


# ── 데이터 ─────────────────────────────────────────────
# 왕: (레인, 시작BC, 끝BC, 이름, 연대표기, 성경, 추정?)
KINGS = [
    ("egypt", 610, 595, "느고 2세", "610–595", "왕하 23:29-35 · 렘 46:2", False),
    ("egypt", 595, 589, "삼메티쿠스 2세", "595–589", "(성경 언급 없음)", False),
    ("egypt", 589, 570, "호브라 (아프리에스)", "589–570", "렘 37:5 · 44:30", False),
    ("egypt", 570, 526, "아마시스", "570–526", "(성경 언급 없음) · 568 느부갓네살 애굽 침공(렘 43:10-13)", False),
    ("judah", 640, 609, "요시야", "640–609", "왕하 22–23 · 대하 34–35", False),
    ("judah", 609, 608.75, "여호아하스", "609 (3개월)", "왕하 23:31-34", False),
    ("judah", 608.75, 598, "여호야김", "609–598", "왕하 23:34–24:6 · 렘 36", False),
    ("judah", 598, 597, "여호야긴", "598–597 (3개월)", "왕하 24:8-15", False),
    ("judah", 597, 586, "시드기야", "597–586", "왕하 24:17–25:7 · 렘 37–39", False),
    ("gov", 586, 582, "그달랴 (총독)", "586–582?", "왕하 25:22-25 · 렘 40–41", True),
    ("gov", 538, 536, "세스바살 (총독)", "538", "스 1:8 · 5:14", True),
    ("gov", 536, 515, "스룹바벨 (총독)", "538?–515?", "스 2:2 · 학 1:1", True),
    ("gov", 445, 433, "느헤미야 (총독)", "445–433", "느 5:14 · 13:6", False),
    ("bab", 626, 605, "나보폴라살", "626–605", "(성경 언급 없음) · 612 니느웨 함락", False),
    ("bab", 605, 562, "느부갓네살 2세", "605–562", "왕하 24–25 · 렘 · 단 1–4", False),
    ("bab", 562, 560, "에윌므로닥", "562–560", "왕하 25:27 · 렘 52:31", False),
    ("bab", 560, 556, "네리글리살", "560–556", "렘 39:3 '네르갈사레셀'?", True),
    ("bab", 556, 555.6, "라바시마르둑", "556 (수개월)", "(성경 언급 없음)", False),
    ("bab", 555.6, 539, "나보니두스", "556–539", "(성경 언급 없음) · 데마에 10년 체류", False),
    ("co", 553, 539, "벨사살 (공동 통치)", "c.553–539", "단 5 · 7:1 · 8:1", True),
    ("co", 539, 537, "메대 다리오", "539–?", "단 5:31 · 6 · 9:1 — 정체 논쟁", True),
    ("persia", 559, 530, "고레스 2세", "559–530 (539 바벨론 정복)", "대하 36:22-23 · 스 1 · 사 44:28; 45:1", False),
    ("persia", 530, 522, "캄비세스 2세", "530–522", "(성경 언급 없음) · 525 애굽 정복", False),
    ("persia", 522, 521.6, "바르디야(스메르디스)", "522 (수개월)", "(성경 언급 없음)", False),
    ("persia", 521.6, 486, "다리오 1세", "522–486", "스 4:24–6:15 · 학 1:1 · 슥 1:1", False),
    ("persia", 486, 465, "크세르크세스 1세 = 아하수에로", "486–465", "에 1–10 · 스 4:6", False),
    ("persia", 465, 424, "아닥사스다 1세", "465–424", "스 4:7-23 · 7–8 · 느 2 · 13:6", False),
    ("persia", 423, 404, "다리오 2세", "423–404", "느 12:22? (논쟁)", True),
]
LANES = [("egypt", "애굽", "바로"), ("judah", "유다", "왕"), ("gov", "유다", "총독"), ("bab", "바벨론", "신바벨론 제국"),
         ("co", "바벨론·메대", "공동 통치 · 논쟁"), ("persia", "메대·바사", "아케메네스 제국")]
LANE_COLOR = {"judah": "judah", "gov": "gov"}
# 여호야긴의 포로 37년 (가는 점선 막대)
JEHOIACHIN_EXILE = (597, 561, "여호야긴 포로 37년 → 561 석방, 왕의 상에서 먹음 (왕하 25:27-30)")

ERAS = [
    (630, 586, "유다 왕국의 마지막 — 경고의 시대", RGB("3A2222")),
    (586, 538, "포로기 — 바벨론의 70년", RGB("3A2E1E")),
    (538, 515, "1차 귀환 · 성전 재건", RGB("1E3326")),
    (515, 458, "58년의 공백 — 에스더", RGB("2A2A2A")),
    (458, 420, "말씀과 성벽의 회복 — 에스라 · 느헤미야 · 말라기", RGB("1E2A3A")),
]

SEVENTY = [
    (605, 538, "약 70년 ① — 바벨론을 섬김 (605 → 539/538)", "렘 25:11-12 · 29:10 · 대하 36:21 · 단 9:2"),
    (586, 515, "약 70년 ② — 성전의 황폐 (586 → 515)", "슥 1:12 '칠십 년 동안 노하셨나이다' · 슥 7:5"),
]

# 사건: (BC, 제목, 성경, 종류, 핵심?)
EVENTS = [
    (627, "예레미야 소명 (요시야 13년)", "렘 1:2", "word", False),
    (622, "율법책 발견 · 요시야 개혁 · 여선지자 훌다", "왕하 22:8, 14 · 23", "word", False),
    (612, "니느웨 함락 — 나훔 예언 성취", "나 3:7", "pol", False),
    (609, "요시야 므깃도 전사", "왕하 23:29", "judg", False),
    (605, "갈그미스 전투 · 1차 포로 (다니엘)", "렘 46:2 · 단 1:1", "judg", False),
    (597, "2차 포로: 여호야긴 · 에스겔 · 1만 명", "왕하 24:14 · 겔 1:2", "judg", False),
    (588, "예루살렘 포위 시작 (18개월)", "왕하 25:1", "judg", False),
    (586, "성전·성벽 불탐 · 대제사장 스라야 처형", "왕하 25:9, 18-21", "judg", True),
    (582, "3차 포로 745명 (합계 4,600)", "렘 52:30", "judg", False),
    (568, "느부갓네살 애굽 침공 (추정)", "렘 43:10-13 · 겔 29:19", "pol", False),
    (561, "여호야긴 석방", "왕하 25:27", "ret", False),
    (539, "바벨론 함락 — 벨사살의 마지막 밤", "단 5:30-31", "pol", True),
    (538, "고레스 칙령 · 1차 귀환 42,360명", "대하 36:22-23 · 스 1–2", "ret", True),
    (536, "성전 기초 — 울음과 기쁨", "스 3:8-13", "ret", False),
    (530, "공사 중단 (536→520, 약 16년)", "스 4:1-5, 24", "judg", False),
    (525, "캄비세스, 애굽 정복", "", "pol", False),
    (520, "공사 재개 — 학개 · 스가랴", "스 5:1-2 · 학 1:14-15", "ret", False),
    (515, "제2성전 완공 · 유월절 (다리오 6년)", "스 6:15-22", "ret", True),
    (483, "아하수에로의 잔치 (3년)", "에 1:3", "pol", False),
    (479, "에스더 왕후 (7년)", "에 2:16-17", "pol", False),
    (474, "하만의 조서 → 부림절 (12년)", "에 3:7 · 9:17-26", "ret", False),
    (452, "성벽 공사 중지 조서 (연대 추정)", "스 4:7-23", "judg", False),
    (458, "에스라 귀환 (4개월 여정) → 통혼 회개", "스 7:7-9 · 9–10", "word", True),
    (445, "느헤미야 귀환 · 성벽 52일", "느 2:1 · 6:15", "ret", True),
    (444.5, "수문 앞 율법 낭독 · 초막절", "느 8", "word", False),
    (433, "느헤미야 바사 복귀 → 재귀환 · 개혁", "느 13:6-7", "word", False),
]

# 선지자·지도자: (번호, 종류, 이름, [(시작, 끝, 확정?)], 막대 표기, 카드)
P = "prophet"; L = "leader"
FIGS = [
    (1, P, "예레미야", [(627, 586, True), (586, 580, False)], "627–586 → 애굽(580?)", {
        "활동": "BC 627–586 이후 (요시야 13년 ~ 애굽 다바네스) · 약 45년, 유다 왕 5명",
        "대상": "유다 왕과 예루살렘 · 바벨론 포로민(29장 편지) · 애굽 피난민(43–44장)",
        "핵심 메시지": "\"돌아오라\"(3:12) → \"바벨론 왕을 섬기라 그리하면 살리라\"(27:12) → \"70년이 차면 돌아오게 하리라\"(29:10) → 새 언약(31:31-33)",
        "갈등": "여호야김 — 두루마리를 칼로 베어 태움(36:23) · 거짓 선지자 하나냐 — 나무 멍에를 꺾음(28장) · 제사장 바스훌 — 차꼬(20:2) · 고관들 — 진흙 구덩이(38:6) · 시드기야 — 사람을 두려워함(38:19) · 요하난 — 애굽행(43장)",
        "대표 본문": "렘 1:4-10 · 7:1-15 · 29:10-14 · 31:31-34 · 36 · 38",
        "하나님의 의도": "심판 직전까지 살 길을 열어 두심 — \"평안이요 재앙이 아니니라\"(29:11)"}),
    (2, P, "스바냐", [(630, 622, False)], "요시야 시대 (개혁 전?)", {
        "활동": "요시야 시대(습 1:1) · 개혁(622) 이전으로 보는 견해가 많음",
        "대상": "유다와 예루살렘 · 블레셋·모압·암몬·구스·앗수르",
        "핵심 메시지": "여호와의 큰 날(1:14) · 겸손한 남은 자(3:12) · \"너로 말미암아 기쁨을 이기지 못하시며\"(3:17)",
        "갈등": "바알의 남은 것과 그마림(1:4) · 말감으로 맹세하는 자(1:5) · \"여호와는 복도 화도 내리지 않으신다\" 하는 자(1:12)",
        "대표 본문": "습 1:4-6, 12-14 · 2:3 · 3:12-17",
        "하나님의 의도": "심판의 경고로 개혁을 깨우시고, 남은 자를 노래하며 기뻐하심"}),
    (3, P, "하박국", [(608, 598, False)], "c.608–598", {
        "활동": "연대 미상 · \"갈대아 사람을 일으켰나니\"(1:6) → 여호야김 시대(608–598) 추정",
        "대상": "하나님께 드리는 질문과 대답 (대화체)",
        "핵심 메시지": "\"의인은 그의 믿음으로 말미암아 살리라\"(2:4) · \"무화과나무가 무성하지 못하며… 나는 여호와로 말미암아 즐거워하리로다\"(3:17-18)",
        "갈등": "선지자 ↔ 하나님 — \"어찌하여 내게 죄악을 보게 하시며\"(1:3) · \"어찌하여 악인이 자기보다 의로운 사람을 삼키는데 잠잠하시나이까\"(1:13)",
        "대표 본문": "합 1:2-4, 12-13 · 2:1-4 · 3:17-19",
        "하나님의 의도": "이해되지 않는 때에도 하나님을 신뢰하며 사는 법을 가르치심"}),
    (4, P, "다니엘", [(605, 536, True)], "605–536 (70년 전체)", {
        "활동": "605 (여호야김 3년, 단 1:1) – 536 (고레스 3년, 10:1) · 포로 70년 전체",
        "대상": "바벨론·메대·바사 왕궁 · 포로민",
        "핵심 메시지": "\"왕들을 폐하시고 왕들을 세우시며\"(2:21) · 뜻을 정함(1:8) · 예레미야의 70년을 읽고 기도(9:2-19) · 70이레(9:24)",
        "갈등": "느부갓네살 — 금 신상, 세 친구의 풀무불(3장) · 벨사살 — 성전 그릇으로 잔치(5:2) · 총리·고관들 — 사자 굴(6장)",
        "대표 본문": "단 1:8 · 2:21 · 3:17-18 · 5:25-28 · 6:10 · 9:2-19",
        "하나님의 의도": "성전이 없는 땅에서도 하나님이 역사의 왕이심을 보이심"}),
    (5, P, "에스겔", [(593, 571, True)], "593–571", {
        "활동": "593 (여호야긴 포로 5년, 겔 1:2) – 571 (27년, 29:17)",
        "대상": "바벨론 그발 강가의 포로민 (597년 2차 포로와 함께 끌려감)",
        "핵심 메시지": "영광이 성전을 떠남(10–11장) · 각자 자기 죄로(18장) · 새 마음과 새 영(36:26) · 마른 뼈(37장) · 영광의 귀환(43:2-5)",
        "갈등": "\"곧 돌아간다\"는 포로민의 낙관 · 거짓 선지자들(13장) · 우상을 마음에 들인 장로들(14:3) · \"아버지가 신 포도를 먹었으므로…\" 속담(18:2)",
        "대표 본문": "겔 1 · 8–11 · 18 · 36:24-28 · 37:1-14 · 43:1-7",
        "하나님의 의도": "무너뜨리신 뒤 새 마음으로 다시 세우심 — \"여호와삼마\"(48:35)"}),
    (6, P, "오바댜", [(586, 550, False)], "연대 미상 (586 이후?)", {
        "활동": "연대 미상 · 예루살렘 함락(586) 이후로 보는 견해가 다수",
        "대상": "에돔",
        "핵심 메시지": "형제의 환난 날에 방관하고 약탈한 에돔의 심판(10-14) · \"나라가 여호와께 속하리라\"(21)",
        "갈등": "에돔(에서) ↔ 유다(야곱) — 형제의 배신 (시 137:7과 같은 장면)",
        "대표 본문": "옵 3-4 · 10-15 · 17 · 21",
        "하나님의 의도": "교만한 자를 낮추시고 시온을 회복하심"}),
    (7, P, "학개", [(520, 519.6, True)], "520 · 4개월", {
        "활동": "520 · 다리오 2년 6월 1일 ~ 9월 24일 (학 1:1; 2:20) — 약 4개월",
        "대상": "총독 스룹바벨 · 대제사장 여호수아 · 남은 백성",
        "핵심 메시지": "\"이 성전이 황폐하였거늘 너희가 판벽한 집에 거주하는 것이 옳으냐\"(1:4) · \"내가 너희와 함께 하노라\"(1:13; 2:4) · 나중 영광이 이전 영광보다 크리라(2:9)",
        "갈등": "\"여호와의 전을 건축할 시기가 이르지 아니하였다\"(1:2) 하는 백성 — 성전 vs 생계",
        "대표 본문": "학 1:2-15 · 2:4-9 · 2:23",
        "하나님의 의도": "16년 멈춘 공사를 다시 일으켜 함께 계심을 회복하심 — 같은 달 24일에 공사 시작(1:15)"}),
    (8, P, "스가랴", [(520, 518, True), (518, 480, False)], "520–518 (+9–14장 연대 미상)", {
        "활동": "520 (다리오 2년 8월, 슥 1:1) – 518 (4년, 7:1) · 9–14장은 연대 미상",
        "대상": "귀환 공동체 · 스룹바벨과 여호수아",
        "핵심 메시지": "\"내게로 돌아오라… 내가 너희에게로 돌아가리라\"(1:3) · 밤의 환상 8개 · \"힘으로 되지 아니하며… 오직 나의 영으로\"(4:6) · 겸손하여 나귀를 타신 왕(9:9)",
        "갈등": "작은 일의 날을 멸시하는 낙심(4:10) · 대제사장을 고발하는 사탄(3:1) · 형식만 남은 금식(7:3-7)",
        "대표 본문": "슥 1:3, 12-17 · 3:1-5 · 4:6-10 · 8:3-8 · 9:9",
        "하나님의 의도": "낙심한 공동체에 성전 너머 메시아의 소망을 심으심"}),
    (9, P, "말라기", [(460, 430, False)], "c.460–430 (추정)", {
        "활동": "연대 미상 · 성전 운영 중, '총독'(1:8), 느헤미야 13장과 같은 문제 → c.460–430 추정",
        "대상": "제사장들 · 백성",
        "핵심 메시지": "\"내가 너희를 사랑하였노라\"(1:2) · 흠 있는 제물(1:8) · 언약의 아내를 버림(2:14-16) · 온전한 십일조(3:10) · 엘리야를 보내리라(4:5-6)",
        "갈등": "제사장·백성의 반문 \"어떻게…?\"(1:2, 6, 7 · 2:14, 17 · 3:7, 8, 13) — 무뎌진 마음",
        "대표 본문": "말 1:2, 6-8 · 2:14-16 · 3:1, 7-10 · 4:2-6",
        "하나님의 의도": "무뎌진 사랑을 깨우고 오실 이를 기다리게 하심 — 이후 약 400년의 침묵"}),
    (10, L, "세스바살 · 스룹바벨 · 여호수아", [(538, 515, True)], "538–515", {
        "활동": "538 귀환 (스 1:8, 11 · 2:2) ~ 515 성전 완공 (6:15)",
        "대상": "1차 귀환자 42,360명 (스 2:64)",
        "핵심 메시지": "제단 먼저(3:2-3) · 성전 기초 — 노인들은 울고 젊은이들은 기뻐함(3:12-13) · \"하나님의 선지자들이 함께 있어 도왔더라\"(5:2)",
        "갈등": "\"그 땅 백성\"의 방해·고발(4:1-5, 24) · 총독 닷드내의 조사(5:3-17) → 다리오가 고레스 조서를 찾아 줌(6:1-12)",
        "대표 본문": "스 1–6 · 학 1–2 · 슥 3–4",
        "하나님의 의도": "건물의 회복 — 예배의 자리를 다시 세우심"}),
    (11, L, "에스더 · 모르드개", [(483, 473, True)], "483–473", {
        "활동": "아하수에로 3년 (483, 에 1:3) ~ 12년 (474, 3:7) → 부림절",
        "대상": "바사 제국 안의 유다 민족 (귀환하지 않은 다수)",
        "핵심 메시지": "\"이때를 위함이 아닌지 누가 알겠느냐\"(4:14) · \"죽으면 죽으리이다\"(4:16)",
        "갈등": "아각 사람 하만(3:1) ↔ 무릎 꿇지 않은 모르드개(3:2) — 민족 말살 조서",
        "대표 본문": "에 3–4 · 7 · 9:20-28",
        "하나님의 의도": "하나님 이름이 한 번도 나오지 않는 책 — 숨은 섭리로 백성을 지키심 (성전과 에스라 사이 58년)"}),
    (12, L, "에스라", [(458, 440, True)], "458–440?", {
        "활동": "458 (아닥사스다 7년, 스 7:7) ~ 445 이후 (느 8 · 12:36)",
        "대상": "2차 귀환자 (남자 약 1,500명, 스 8) · 예루살렘 공동체",
        "핵심 메시지": "연구하고 → 준행하며 → 가르치기로 결심(7:10) · \"하나님의 선한 손\"(7:6, 9 · 8:22)",
        "갈등": "이방 통혼한 방백·제사장·레위인(9:1-2 · 10:18) — 에스라는 \"우리\"로 회개(9:6) · 반대한 요나단·야스야(10:15)",
        "대표 본문": "스 7:6-10 · 8:21-23 · 9:3-15 · 느 8:1-12",
        "하나님의 의도": "말씀으로 사람을 회복하심 — \"여호와로 인하여 기뻐하는 것이 너희의 힘\"(느 8:10)"}),
    (13, L, "느헤미야", [(445, 433, True), (433, 428, False)], "445–433 → 재귀환", {
        "활동": "445 (아닥사스다 20년, 느 2:1) ~ 433 (32년, 13:6) → 얼마 후 재귀환",
        "대상": "예루살렘 주민 · 귀족 · 제사장",
        "핵심 메시지": "금식 기도(1:4-11) · 성벽 52일 완공(6:15) · 언약 갱신 서명(9:38–10장) · 안식일·십일조 개혁(13장)",
        "갈등": "산발랏·도비야·게셈의 조롱과 음모(2:19 · 4 · 6) · 고리대 하는 귀족(5:1-13) · 성전 방을 도비야에게 내준 엘리아십(13:4-9)",
        "대표 본문": "느 1:4-11 · 2:17-20 · 4:14 · 6:15-16 · 13",
        "하나님의 의도": "공동체(울타리)의 회복 — \"내 하나님의 선한 손이 나를 도우심\"(2:8, 18)"}),
]

HEART = [
    (630, 586, "경고 — \"돌아오라\"", "선지자를 끝까지 보내심"),
    (586, 538, "징계 — 그러나 70년 약속", "본심이 아니시로다 (애 3:33)"),
    (538, 520, "귀환", "제단과 기초"),
    (520, 515, "예배 회복", "성전 완공"),
    (515, 458, "숨은 섭리", "에스더 — 이름 없이 지키심"),
    (458, 445, "말씀 회복", "에스라"),
    (445, 430, "공동체 회복", "느헤미야"),
    (430, 420, "메시아 대망", "말라기 → 400년"),
]

NOTES = [
    ("연대 기준", "왕들의 연대는 바벨론 연대기 · 페르시아 왕명록에 따른 관례적 연대. 예루살렘 함락은 586년설(587년설도 있음). 유다 왕 연대는 Thiele 계열."),
    ("메대 다리오 (단 5:31; 6; 9:1)", "성경 밖 기록에서 이름이 확인되지 않아 논쟁 — ① 바벨론 총독 구바루(고브리아스)설 ② 고레스 동일인설(단 6:28을 '다리오 시대 곧 고레스 시대'로 읽음) ③ 기타. 연표에는 점선으로만 표시."),
    ("벨사살", "나보니두스의 아들. 나보니두스가 아라비아 데마에 머문 약 10년(c.553–543) 동안 바벨론을 맡은 공동 통치자 — 그래서 다니엘에게 '셋째 통치자'(단 5:29)를 약속."),
    ("70년의 계산", "① 605(1차 포로) → 539/538(고레스) ② 586(성전 파괴) → 515(성전 완공) ③ 상징적 완전수. 성경 안에서도 렘 25·29(바벨론의 70년)와 슥 1:12·7:5(진노·금식의 70년)가 다른 시점을 센다."),
    ("스 4:6-23", "아하수에로·아닥사스다 때의 고발(성벽 공사 중지)을 성전 재건 이야기 중간에 주제별로 모아 둔 단락 — 시간 순서로는 486년 이후. 성벽 중지 조서는 c.460–446 사이로 추정(연표에는 452 근처)."),
    ("에스라의 연대", "아닥사스다 1세 7년 = 458년(전통적 견해, 이 연표). 아닥사스다 2세(398년)로 보는 소수설이 있음."),
    ("연대 미상 선지서", "스바냐(요시야 시대, 개혁 전 추정) · 하박국(여호야김 시대 추정) · 오바댜(586 이후 다수설) · 스가랴 9–14장 · 말라기(c.460–430)는 본문에 연대가 없어 점선 처리."),
    ("네리글리살", "렘 39:3, 13의 '네르갈사레셀'(바벨론 고관)과 같은 인물일 가능성 — 확정은 아님."),
    ("훌다·나훔", "여선지자 훌다(왕하 22:14, 622)와 나훔(니느웨 함락 612 이전)은 사건 레이어에만 표시."),
]

# ── 그리기 도구 ─────────────────────────────────────────
prs = Presentation()
SH_DEFAULT = 31.5
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH_DEFAULT)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _font(r, size, color, bold):
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        rPr.append(rPr.makeelement(NS + tag, {"typeface": FONT}))


def box(x, y, w, h, fill=None, line=None, dash=False, shape=MSO_SHAPE.RECTANGLE, lw=1.0, adj=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(max(w, 0.01)), Inches(h))
    if adj is not None:
        sp.adjustments[0] = adj
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = MSO_LINE.DASH
    sp.shadow.inherit = False
    return sp


def text(x, y, w, h, paras, size=12, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=False):
    """paras: 문자열, 또는 [[(글, 크기, 색, 굵게), ...], ...] (문단별 run 목록)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [[(ln, size, color, bold)] for ln in paras.split("\n")]
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, sz, c, b in runs:
            r = p.add_run(); r.text = t; _font(r, sz, c, b)
    return tb


def line(x0, y0, x1, y1, color, w=0.75, dash=False):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    ln.line.color.rgb = color; ln.line.width = Pt(w)
    if dash:
        ln.line.dash_style = MSO_LINE.DASH
    return ln


def tw(t, size):
    """대략적인 글자 폭(인치)."""
    em = size / 72
    return sum(em * (0.55 if ord(c) < 0x1100 else 0.95) for c in t)


def place_rows(items, gap=0.15):
    """items: (x_left, width) 목록 → 겹치지 않는 행 번호 목록 (탐욕 배치)."""
    ends, rows = [], []
    for xl, w in items:
        for r, e in enumerate(ends):
            if xl >= e + gap:
                ends[r] = xl + w; rows.append(r); break
        else:
            ends.append(xl + w); rows.append(len(ends) - 1)
    return rows, len(ends)


def grid_lines(y0, y1):
    for bc in range(Y_START, Y_END - 1, -10):
        line(X(bc), y0, X(bc), y1, DIM, 0.5, dash=True)
    for bc, col in [(586, COL["judg"]), (539, GREY), (538, COL["ret"]), (515, COL["ret"]), (458, COL["word"]), (445, COL["ret"])]:
        line(X(bc), y0, X(bc), y1, col, 1.25, dash=True)


def lane_label(y, h, title, sub, color):
    box(0.25, y, LBL - 0.45, h, fill=RGB("222222"), line=color, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
    text(0.35, y, LBL - 0.65, h, [[(title, 15, color, True)], [(sub, 10, GREY, False)]], anchor=MSO_ANCHOR.MIDDLE)


# ── 1) 제목·범례 ────────────────────────────────────────
y = 0.3
box(0.3, y, 21.5, 1.0, fill=RGB("6B4F2A"), shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
text(0.6, y, 21.0, 1.0, "포로기와 귀환 — 왕들과 선지자들 (BC 630 – 420)", 34, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(22.2, y + 0.05, 33, 0.5,
     "마스터 연표(솔로몬 → BC 586) 이어서 · 네모의 왼쪽 변 = 즉위/활동 시작, 오른쪽 변 = 폐위/활동 끝 · 점선 = 연대 추정·논쟁 · ①~⑬ = 아래 카드 번호", 14, GREY)
leg = [("심판·멸망", COL["judg"]), ("귀환·회복", COL["ret"]), ("말씀 사건", COL["word"]), ("정치 사건", COL["pol"]),
       ("선지자", COL["prophet"]), ("말씀 지도자", COL["leader"]), ("애굽", COL["egypt"]), ("유다", COL["judah"]),
       ("바벨론", COL["bab"]), ("메대·바사", COL["persia"])]
lx = 22.2
for name, c in leg:
    box(lx, y + 0.68, 0.32, 0.22, fill=c)
    text(lx + 0.4, y + 0.62, 2.5, 0.32, name, 13, GREY)
    lx += 0.55 + tw(name, 13) + 0.35
y += 1.35

# ── 2) 시간축 ──────────────────────────────────────────
AX_Y = y
text(0.3, AX_Y, LBL - 0.4, 0.5, "BC", 16, GREY, True)
for bc in range(Y_START, Y_END - 1, -5):
    major = bc % 10 == 0
    line(X(bc), AX_Y + (0.32 if major else 0.42), X(bc), AX_Y + 0.55, GREY if major else DIM, 1.0)
    if major:
        text(X(bc) - 0.5, AX_Y, 1.0, 0.3, str(bc), 14, WHITE, True, PP_ALIGN.CENTER)
line(X(Y_START), AX_Y + 0.55, X(Y_END), AX_Y + 0.55, GREY, 1.5)
y = AX_Y + 0.7
GRID_TOP = y

# ── 3) 시대 띠 ─────────────────────────────────────────
for a, b, name, fill in ERAS:
    box(X(a), y, X(b) - X(a), 0.5, fill=fill, line=DIM, lw=0.5)
    text(X(a) + 0.1, y, X(b) - X(a) - 0.2, 0.5, name, 14 if X(b) - X(a) > 6 else 11, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
y += 0.65

# ── 4) 왕 레인 ─────────────────────────────────────────
BOX_H = 0.95
lane_y = {}
for key, title, sub in LANES:
    ks = [k for k in KINGS if k[0] == key]
    # 칸이 좁아 이름이 안 들어가는 왕은 아래 줄에 따로 표시
    outs = []
    for k in ks:
        w = X(k[2]) - X(k[1])
        need = max(tw(k[3], 14), tw(k[4], 11), tw(k[5], 10)) + 0.2
        if w < need:
            outs.append(k)
    items = [(X(k[1]) + (X(k[2]) - X(k[1])) / 2 - (max(tw(k[3], 13), tw(k[4] + " · " + k[5], 10)) + 0.1) / 2,
              max(tw(k[3], 13), tw(k[4] + " · " + k[5], 10)) + 0.1) for k in outs]
    orows, n_orow = place_rows(items, 0.2) if outs else ([], 0)
    extra = 0.15 + n_orow * 0.62 if outs else 0
    if key == "judah":
        extra = max(extra, 0.0) + 0.4  # 여호야긴 포로 37년 막대
    lane_h = BOX_H + extra + 0.1
    lane_y[key] = y
    color = COL[LANE_COLOR.get(key, key)]
    lane_label(y, lane_h, title, sub, color)
    line(LBL, y + lane_h + 0.04, LBL + W, y + lane_h + 0.04, RGB("2E2E2E"), 0.75)
    for k in ks:
        _, a, b, name, yrs, ref, est = k
        x0, x1 = X(a), X(b)
        w = max(x1 - x0, 0.06)
        sp = box(x0, y + 0.05, w, BOX_H - 0.1, fill=color, line=WHITE if not est else GREY, dash=est, lw=1.0)
        if est:
            sp.fill.fore_color.rgb = RGB("%02X%02X%02X" % tuple(int(int(str(color)[i:i + 2], 16) * 0.55) for i in (0, 2, 4)))
        if k not in outs:
            text(x0 + 0.1, y + 0.08, w - 0.15, BOX_H - 0.15,
                 [[(name, 14, INK if not est else WHITE, True)], [(yrs, 11, INK if not est else WHITE, False)],
                  [(ref, 10, INK if not est else WHITE, False)]], anchor=MSO_ANCHOR.MIDDLE)
    for k, (xl, wd), r in zip(outs, items, orows):
        _, a, b, name, yrs, ref, est = k
        ly = y + BOX_H + 0.1 + r * 0.62
        cx = X(a) + (X(b) - X(a)) / 2
        line(cx, y + BOX_H - 0.05, cx, ly, color, 0.75)
        text(xl, ly, wd, 0.6, [[(name, 13, color, True)], [(yrs + " · " + ref, 10, WHITE, False)]])
    if key == "judah":
        a, b, label = JEHOIACHIN_EXILE
        by = y + lane_h - 0.42
        box(X(a), by, X(b) - X(a), 0.14, fill=RGB("2C4A63"), line=COL["judah"], dash=True, lw=0.75)
        text(X(a) + 0.1, by + 0.15, 20, 0.25, label, 10, COL["judah"])
    y += lane_h + 0.08

# ── 5) 70년 괄호 ───────────────────────────────────────
y += 0.1
for i, (a, b, title, ref) in enumerate(SEVENTY):
    yy = y + i * 0.55
    x0, x1 = X(a), X(b)
    line(x0, yy + 0.1, x0, yy + 0.4, COL["word"], 2)
    line(x1, yy + 0.1, x1, yy + 0.4, COL["word"], 2)
    line(x0, yy + 0.25, x1, yy + 0.25, COL["word"], 2)
    text(x1 + 0.2, yy + 0.05, 25, 0.4, [[(title + "  ", 14, COL["word"], True), (ref, 11, GREY, False)]])
lane_label(y, 1.05, "70년", "두 가지 계산", COL["word"])
y += 1.25

# ── 6) 사건 레이어 ─────────────────────────────────────
EV_SIZE, EV_REF = 12, 10
ev_items = []
for bc, title, ref, kind, key in EVENTS:
    wdt = max(tw(title, EV_SIZE + (1 if key else 0)), tw(f"BC {int(bc + 0.5)} · {ref}", EV_REF)) + 0.15
    ev_items.append((X(bc) - 0.05, wdt))
order = sorted(range(len(EVENTS)), key=lambda i: ev_items[i][0])
rows_sorted, n_ev_rows = place_rows([ev_items[i] for i in order], 0.15)
ev_row = {i: r for i, r in zip(order, rows_sorted)}
EV_ROW_H = 0.62
ev_h = 0.45 + n_ev_rows * EV_ROW_H
lane_label(y, ev_h, "주요 사건", "▼ = 연도", WHITE)
base = y + 0.25
line(LBL, base, LBL + W, base, GREY, 1.0)
# 공사 중단 구간 표시
box(X(536), base - 0.09, X(520) - X(536), 0.18, fill=RGB("5A2020"), line=COL["judg"], dash=True, lw=0.75)
for i, (bc, title, ref, kind, key) in enumerate(EVENTS):
    c = COL[kind]
    xm = X(bc)
    line(xm, base, xm, base + 0.25 + ev_row[i] * EV_ROW_H, c, 0.75)
    box(xm - 0.13, base - 0.2, 0.26, 0.22, fill=c, shape=MSO_SHAPE.FLOWCHART_MERGE)
for i, (bc, title, ref, kind, key) in enumerate(EVENTS):
    c = COL[kind]
    ly = base + 0.2 + ev_row[i] * EV_ROW_H
    title_runs = [(title, EV_SIZE + (1 if key else 0), c, True)]
    tb = text(ev_items[i][0], ly, ev_items[i][1] + 0.1, 0.56,
              [title_runs, [(f"BC {int(bc + 0.5)}" + (f" · {ref}" if ref else ""), EV_REF, GREY, False)]])
    tb.fill.solid(); tb.fill.fore_color.rgb = BG
y += ev_h + 0.25

# ── 7) 선지자·지도자 막대 ──────────────────────────────
BAR_ROW = 0.62
fig_top = y
fig_h = len(FIGS) * BAR_ROW + 0.2
box(0.25, fig_top, LBL - 0.45, fig_h / 2 - 0.05, fill=RGB("222222"), line=COL["prophet"], lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08)
text(0.35, fig_top, LBL - 0.65, fig_h / 2 - 0.05, [[("선지자", 15, COL["prophet"], True)], [("활동 기간", 10, GREY, False)]], anchor=MSO_ANCHOR.MIDDLE)
box(0.25, fig_top + fig_h / 2 + 0.05, LBL - 0.45, fig_h / 2 - 0.05, fill=RGB("222222"), line=COL["leader"], lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08)
text(0.35, fig_top + fig_h / 2 + 0.05, LBL - 0.65, fig_h / 2 - 0.05, [[("말씀 지도자", 15, COL["leader"], True)], [("귀환·재건 인물", 10, GREY, False)]], anchor=MSO_ANCHOR.MIDDLE)
CIRC = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬"
for idx, (no, kind, name, spans, label, card) in enumerate(FIGS):
    yy = fig_top + 0.1 + idx * BAR_ROW
    c = COL[kind]
    line(LBL, yy + BAR_ROW - 0.03, LBL + W, yy + BAR_ROW - 0.03, RGB("262626"), 0.5)
    x_first = X(spans[0][0])
    x_last = X(spans[-1][1])
    for a, b, sure in spans:
        w = max(X(b) - X(a), 0.18)
        sp = box(X(a), yy + 0.06, w, BAR_ROW - 0.18, fill=c if sure else None, line=c, dash=not sure, lw=1.5,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.25)
    lab = [(f"{CIRC[no - 1]} {name}", 14, c, True), (f"   {label}", 11, GREY, False)]
    x_solid_end = X(spans[0][1])
    inside = spans[0][2] and (x_solid_end - x_first) > tw(f"{CIRC[no - 1]} {name}   {label}", 13) + 0.3
    if inside:
        lab = [(f"{CIRC[no - 1]} {name}", 14, INK, True), (f"   {label}", 11, INK, False)]
        text(x_first + 0.12, yy, x_solid_end - x_first, BAR_ROW - 0.06, [lab], anchor=MSO_ANCHOR.MIDDLE)
    else:
        text(x_last + 0.15, yy, 14, BAR_ROW - 0.06, [lab], anchor=MSO_ANCHOR.MIDDLE)
y = fig_top + fig_h + 0.15

# ── 8) 하나님의 마음 ───────────────────────────────────
HEART_H = 1.0
lane_label(y, HEART_H + 0.35, "하나님의 마음", "시대마다", RGB("FF9EC4"))
for i, (a, b, title, sub) in enumerate(HEART):
    shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
    shade = [RGB("5A2A2A"), RGB("5A4426"), RGB("2A5A3A"), RGB("2F6B45"), RGB("3A3A3A"), RGB("6B5A20"), RGB("2A4A6B"), RGB("6B2A55")][i]
    box(X(a), y, X(b) - X(a) + 0.25, HEART_H, fill=shade, line=RGB("FF9EC4"), lw=1.0, shape=shp)
    wide = X(b) - X(a) > 3
    text(X(a) + (0.5 if i else 0.2), y + 0.05, X(b) - X(a) - 0.4, HEART_H - 0.1,
         [[(title, 15 if wide else 11, WHITE, True)], [(sub, 11 if wide else 9, RGB("FFD6E8"), False)]],
         anchor=MSO_ANCHOR.MIDDLE, wrap=not wide)
text(LBL, y + HEART_H + 0.05, W, 0.3,
     "징계 → 귀환 → 예배 회복 → (숨은 섭리) → 말씀 회복 → 공동체 회복 → 메시아 대망  ·  \"내가 너희와 함께 하노라\"(학 1:13)  ·  \"내게로 돌아오라 그리하면 내가 너희에게로 돌아가리라\"(슥 1:3)",
     12, RGB("FFD6E8"))
y += HEART_H + 0.5
GRID_BOTTOM = y - 0.4

# ── 9) 카드 ────────────────────────────────────────────
FIELDS = ["활동", "대상", "핵심 메시지", "갈등", "대표 본문", "하나님의 의도"]
COLS_N = 7
GAP = 0.25
CW = (SW - 0.6 - GAP * (COLS_N - 1)) / COLS_N
F_SIZE = 14
LINE_H = F_SIZE / 72 * 1.32


def n_lines(t, width):
    out, cur = 0, 0.0
    for para in [t]:
        cur = 0.0; out += 1
        for ch in para:
            cw = F_SIZE / 72 * (0.55 if ord(ch) < 0x1100 else 0.95)
            if cur + cw > width:
                out += 1; cur = 0.0
            cur += cw
    return out


text(0.3, y, 30, 0.6, [[("선지자 · 말씀 지도자 카드", 22, WHITE, True),
                        ("   — 활동 · 대상 · 핵심 메시지 · 갈등 · 대표 본문 · 하나님의 의도", 14, GREY, False)]])
y += 0.75
for row_start in range(0, len(FIGS), COLS_N):
    row = FIGS[row_start:row_start + COLS_N]
    heights = []
    for fig in row:
        card = fig[5]
        hh = 0.75
        for f in FIELDS:
            hh += n_lines(f"{f}  {card[f]}", CW - 0.4) * LINE_H + 0.1
        heights.append(hh + 0.15)
    ch = max(heights)
    for j, (no, kind, name, spans, label, card) in enumerate(row):
        cx = 0.3 + j * (CW + GAP)
        c = COL[kind]
        box(cx, y, CW, ch, fill=CARD, line=c, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.03)
        box(cx, y, CW, 0.62, fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
        text(cx + 0.2, y, CW - 0.4, 0.62, [[(f"{CIRC[no - 1]} {name}", 17, INK, True), (f"   {label}", 12, INK, False)]],
             anchor=MSO_ANCHOR.MIDDLE)
        paras = []
        for f in FIELDS:
            paras.append([(f + "  ", F_SIZE, c, True), (card[f], F_SIZE, WHITE, False)])
        tb = text(cx + 0.2, y + 0.75, CW - 0.4, ch - 0.85, paras, wrap=True)
        for p in tb.text_frame.paragraphs:
            p.space_after = Pt(5)
    y += ch + 0.3

# 세로 격자선 (시대 띠 ~ 하나님의 마음)
grid_lines(GRID_TOP, GRID_BOTTOM)
# 격자선이 도형 뒤로 가도록 맨 앞으로 이동
sp_tree = s.shapes._spTree
conns = [e for e in list(sp_tree)[-(len(range(Y_START, Y_END - 1, -10)) + 6):]]
for e in conns:
    sp_tree.remove(e); sp_tree.insert(2, e)

text(SW - 12.4, y, 12, 0.35, "자료: 개역개정 · BHS · 바벨론 연대기 · 페르시아 왕명록 | Steve 설교 위키 (scripts/make_postexile_timeline_ppt.py)",
     10, DIM, align=PP_ALIGN.RIGHT)
y += 0.5
prs.slide_height = Inches(min(max(y, 20), 56))

# ── 2장: 연대 근거와 논쟁 메모 ─────────────────────────
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
SH2 = prs.slide_height / 914400
box(0.6, 0.5, 40, 2.2, fill=RGB("6B4F2A"), shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
text(1.2, 0.5, 39, 2.2, "연대 근거와 논쟁 메모 — 점선으로 표시한 이유", 72, bold=True, anchor=MSO_ANCHOR.MIDDLE)
ny = 3.4
half = math.ceil(len(NOTES) / 2)
nh = (SH2 - ny - 0.6) / half
for i, (k, v) in enumerate(NOTES):
    col = i // half
    cx = 0.6 + col * 27.6
    yy = ny + (i % half) * nh
    box(cx, yy, 26.8, nh - 0.4, fill=CARD, line=COL["word"], lw=2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
    tb = text(cx + 0.6, yy + 0.4, 25.6, nh - 1.0, [[(k, 44, COL["word"], True)], [(v, 34, WHITE, False)]], wrap=True)
    tb.text_frame.paragraphs[0].space_after = Pt(10)

prs.save(OUT)
print("saved", OUT, f"slide 56 x {min(max(y, 20), 56):.1f} in")
